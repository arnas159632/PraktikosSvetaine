from flask import Blueprint, render_template, request, send_file
from io import BytesIO
from datetime import datetime
import zipfile
from utils.db import get_db_connection
from docx import Document
from utils.docx_utils import replace_placeholders, extract_placeholders_and_data

main_routes = Blueprint('main', __name__)

@main_routes.route('/')
def form():
    """Render the form page."""
    return render_template('form.html')

@main_routes.route('/generate-documents', methods=['POST'])
def generate_documents():
    """Generate documents based on the form input."""
    fill2_value = request.form.get('fill2')
    date = request.form.get('date')
    faculty = request.form.get('faculty')

    date_obj = datetime.now().date() if not date else datetime.strptime(date, '%Y-%m-%d').date()

    doc1 = Document("docs/ataskaita.docx")
    doc2 = Document("docs/KatedrosAtaskaita.docx")

    replacements = {f"{{{{{field_name}}}}}": [field_value] for field_name, field_value in request.form.items()}

    placeholders_doc1 = extract_placeholders_and_data(doc1)
    placeholders_doc2 = extract_placeholders_and_data(doc2)

    replaced_data_doc1 = replace_placeholders(doc1, replacements)
    replaced_data_doc2 = replace_placeholders(doc2, replacements)

    file_stream_1 = BytesIO()
    file_stream_2 = BytesIO()
    doc1.save(file_stream_1)
    file_stream_1.seek(0)
    doc2.save(file_stream_2)
    file_stream_2.seek(0)

    with get_db_connection() as conn:
        cursor = conn.cursor()

        def save_document(name, date, faculty, data, placeholders, replaced_data):
            faculty_value = faculty if faculty else None

            cursor.execute(
                "INSERT INTO documents (name, date, faculty, data) VALUES (%s, %s, %s, %s)",
                (name, date, faculty_value, data)
            )
            document_id = cursor.lastrowid

            for placeholder in placeholders.keys():
                cursor.execute(
                    "INSERT IGNORE INTO placeholders (placeholder) VALUES (%s)",
                    (placeholder,)
                )

            cursor.execute("SELECT id, placeholder FROM placeholders")
            placeholder_ids = {placeholder: placeholder_id for placeholder_id, placeholder in cursor.fetchall()}

            for placeholder, texts in placeholders.items():
                placeholder_id = placeholder_ids.get(placeholder)
                if placeholder_id:
                    placeholder_value = ' '.join(texts) if texts else ''
                    cursor.execute(
                        "INSERT IGNORE INTO document_placeholders (document_id, placeholder_id, value) VALUES (%s, %s, %s)",
                        (document_id, placeholder_id, placeholder_value)
                    )

            for placeholder, replaced_value in replaced_data.items():
                cursor.execute(
                    "INSERT IGNORE INTO document_replacements (document_id, placeholder, value) VALUES (%s, %s, %s)",
                    (document_id, placeholder, replaced_value)
                )

            conn.commit()
            return document_id

        file_name_1 = f"{fill2_value}_DestytojoAtaskaita.docx" if fill2_value else 'Atskaita_Document.docx'
        file_name_2 = f"{fill2_value}_KatedrosAtaskaita.docx" if fill2_value else 'KatedrosAtaskaita_Document.docx'

        doc1_id = save_document(file_name_1, date_obj, faculty, file_stream_1.read(), placeholders_doc1, replaced_data_doc1)
        doc2_id = save_document(file_name_2, date_obj, faculty, file_stream_2.read(), placeholders_doc2, replaced_data_doc2)

        zip_stream = BytesIO()
        with zipfile.ZipFile(zip_stream, 'w') as zip_file:
            zip_file.writestr(file_name_1, file_stream_1.getvalue())
            zip_file.writestr(file_name_2, file_stream_2.getvalue())
        
        zip_stream.seek(0)

        return send_file(zip_stream, as_attachment=True, download_name='Documents.zip', mimetype='application/zip')
